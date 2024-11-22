import re
from docx import Document
import openpyxl

def count_keywords_in_docx(docx_path, regex_pattern, xlsx_path, unique_xlsx_path):
    # 编译正则表达式模式
    pattern = re.compile(regex_pattern)

    # 打开DOCX文件
    doc = Document(docx_path)

    # 创建一个新的Excel工作簿并选择活动工作表
    wb = openpyxl.Workbook()
    ws = wb.active

    # 设置Excel文件的标题
    ws.append(["匹配到的关键字"])

    # 初始化关键字计数器和关键字集合
    keyword_count = 0
    keywords_set = set()

    # 遍历文档中的所有段落
    for para in doc.paragraphs:
        text = para.text.strip()
        # 查找段落文本中的所有匹配项
        matches = pattern.findall(text)
        # 将匹配项添加到Excel文件并更新计数
        for match in matches:
            ws.append([match])
            keyword_count += 1
            keywords_set.add(match)

    # 保存Excel文件
    wb.save(xlsx_path)

    # 检查是否有重复项
    has_duplicates = keyword_count != len(keywords_set)
    print(f"总共找到的关键字数量: {keyword_count}")
    print(f"是否有重复项: {'是' if has_duplicates else '否'}")

    # 如果有重复项，生成去掉重复项的输出文件
    if has_duplicates:
        unique_wb = openpyxl.Workbook()
        unique_ws = unique_wb.active
        unique_ws.append(["匹配到的关键字（去重）"])
        for keyword in keywords_set:
            unique_ws.append([keyword])
        unique_wb.save(unique_xlsx_path)

    return keyword_count

# 示例用法
docx_path = './形式化需求.docx'
regex_pattern = r'\b[A-Z][-_]\d+\b'
xlsx_path = './形式化需求.xlsx'
unique_xlsx_path = './形式化需求_去重.xlsx'
count = count_keywords_in_docx(docx_path, regex_pattern, xlsx_path, unique_xlsx_path)
print(f"总共找到的关键字数量: {count}")